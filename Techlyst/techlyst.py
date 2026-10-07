import os
import json
import time
import html
import tempfile
import threading
import traceback
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime
import win32com.client as win32


CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "teclist_config.json")

def carregar_config():
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r") as f:
                return json.load(f)
    except:
        pass
    return {"ultimos_skus": "20104418\n", "pasta_destino": "", "centro": "1000 - Plant-A"}

def salvar_config(dados):
    try:
        cfg = carregar_config()
        cfg.update(dados)
        with open(CONFIG_FILE, "w") as f:
            json.dump(cfg, f)
    except:
        pass


DB_INSUMOS_FABRICA = {
    # [REVIEW] Original dictionary contained ~200+ real material codes and detailed descriptions.
    # Replaced entirely with 10 generic example entries to demonstrate the structure.
    "10000001": "MATERIAL-A",
    "DESC_MATERIAL_A_GENERICO": "MATERIAL-A",
    "10000002": "MATERIAL-B",
    "DESC_MATERIAL_B_GENERICO": "MATERIAL-B",
    "10000003": "MATERIAL-C",
    "DESC_MATERIAL_C_GENERICO": "MATERIAL-C",
    "10000004": "MATERIAL-D",
    "DESC_MATERIAL_D_GENERICO": "MATERIAL-D",
    "10000005": "MATERIAL-E",
    "DESC_MATERIAL_E_GENERICO": "MATERIAL-E"}



def get_db_path():
    appdata = os.environ.get('LOCALAPPDATA') or os.environ.get('APPDATA') or os.path.expanduser('~')
    app_dir = os.path.join(appdata, 'TecListForge')
    os.makedirs(app_dir, exist_ok=True)
    return os.path.join(app_dir, 'teclist_insumos_custom.json')

def carregar_db_hibrido():
    caminho = get_db_path()
    try:
        if os.path.exists(caminho):
            with open(caminho, "r", encoding="utf-8") as f:
                return json.load(f)
    except:
        pass
    
    # Se nao existe ou falhou, retorna o de fabrica
    return {
        "_metadata": {
            "versao": "1.0",
            "data_atualizacao": "Fábrica"
        },
        "data": DB_INSUMOS_FABRICA.copy()
    }

def salvar_db_hibrido(db_dict):
    caminho = get_db_path()
    try:
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(db_dict, f, indent=4, ensure_ascii=False)
        return True
    except Exception as e:
        registrar_log(f"Erro ao salvar DB: {e}", "erro")
        return False

DB_COMPLETO = carregar_db_hibrido()
DB_INSUMOS = DB_COMPLETO.get("data", {})
DB_METADATA = DB_COMPLETO.get("_metadata", {"versao": "1.0", "data_atualizacao": "Fábrica"})

def classificar_insumo(codigo, descricao):
    """Busca o Tipo Insumo pelo código ou descrição no banco de dados JSON."""
    if codigo and codigo in DB_INSUMOS:
        return DB_INSUMOS[codigo]
    if descricao and descricao in DB_INSUMOS:
        return DB_INSUMOS[descricao]
    return "NÃO CLASSIFICADO"


_stop_flag        = threading.Event()   # sinaliza cancelamento
_tempo_inicio     = None                # datetime do início

# ── Controle por fase ──────────────────────────────────────────────
_fase_atual       = ""                  # descrição da fase ativa
_fase_inicio      = None                # quando a fase atual começou
_fase_concluidos  = 0                   # unidades concluídas na fase atual
_fase_total       = 0                   # total de unidades da fase atual
_timer_ativo      = False               # controla o loop do timer

_total_zs_encontrados = 0               # rastreia a qtde de Zs para o ETA determinístico

_historico_tempos = []                  # lista de segundos gastos por unidade


BG_ROOT      = "#0b0f19"
BG_HEADER    = "#111827"
BG_CARD      = "#1f2937"
BG_FOOTER    = "#0f1422"
ACCENT       = "#3b82f6"
ACCENT_DIM   = "#2563eb"
TEXT_PRIMARY = "#f3f4f6"
TEXT_MUTED   = "#9ca3af"
TEXT_WARN    = "#fbbf24"
TEXT_ERROR   = "#ef4444"
TEXT_LOG     = "#93c5fd"

FNT_TITLE    = ("Segoe UI", 13, "bold")
FNT_SUB      = ("Segoe UI", 9)
FNT_LOG      = ("Cascadia Code", 9)
FNT_STATUS   = ("Segoe UI", 8)
FNT_LABEL    = ("Segoe UI", 8, "bold")
FNT_BTN      = ("Segoe UI", 10, "bold")


# Machine names updated to match Visconf: PH.04, PH.06, PH.07, PH.08, PH.12..15, DMX.05, FOR.01..03, PROF.01..04
DEPARA_LINHAS = {
    "B-CONVAM": "PH.04",
    "B-CONVBR": "PH.06",
    "B-CONVCR": "PH.07",
    "B-CONVCB": "FOR.01",
    "B-CONVC2": "FOR.02",
    "B-CONVP7": "DMX.05",
    "B-CONVTA": "PROF.01",
    "B-CONVL4": "PROF.02",
    "B-CONVL5": "PH.08",
    "B-CONVL6": "PH.12",
    "B-CONVL7": "PH.13",
    "B-CONVL8": "PH.14",
    "B-CONV12": "PH.15",
    "B-CONV13": "FOR.03",
    "B-CONV14": "PROF.03",
    "B-CONV15": "PROF.04",
}

def calcular_unidade_base(linha_formatada):
    cx = ["PROF.01", "PROF.02", "PROF.03", "PROF.04", "FOR.01", "FOR.02", "FOR.03", "PH.04", "PH.06", "PH.07", "PH.08", "PH.12", "PH.13", "PH.14", "PH.15", "DMX.05"]
    fdo = ["FOR.01", "FOR.02", "FOR.03"]
    kg = ["MP04", "MP07", "MP08", "MP09", "REB07", "AT01"]
    
    if linha_formatada in cx: return "CX"
    if linha_formatada in fdo: return "FDO"
    if linha_formatada in kg: return "KG"
    return "N/A"

def limpar_numero(val_str):
    if not val_str:
        return 0.0
    val_str = str(val_str).strip()
    
    is_negative = False
    if val_str.endswith("-"):
        is_negative = True
        val_str = val_str[:-1]
        
    val_str = val_str.replace(".", "").replace(",", ".")
    try:
        num = float(val_str)
        if is_negative:
            num = -num
        return num
    except ValueError:
        return 0.0


def conectar_sap_ou_abrir():
    """Realiza a verificação, abertura visível e conexão ao SAP."""
    registrar_log("Verificando conexão com o SAP...")
    SAPGui = None
    try:
        SAPGui = win32.GetObject("SAPGUI")
    except Exception:
        pass

    if SAPGui is None:
        registrar_log("SAP fechado. Abrindo o SAP Logon na tela...")
        try:
            caminho_sap = r"C:\Program Files (x86)\SAP\FrontEnd\SAPgui\saplogon.exe"
            if not os.path.exists(caminho_sap):
                caminho_sap = r"C:\Program Files\SAP\FrontEnd\SAPgui\saplogon.exe"
            os.startfile(caminho_sap)
        except Exception:
            try:
                subprocess.Popen([caminho_sap])
            except Exception as ex:
                raise Exception(f"Não foi possível iniciar o SAP Logon: {ex}")
        
        registrar_log("Aguardando o SAP Logon aparecer...")
        for _ in range(30):
            try:
                SAPGui = win32.GetObject("SAPGUI")
                if SAPGui is not None: break
            except: pass
            time.sleep(1)
        else:
            raise Exception("Tempo limite (30s) excedido aguardando o SAP Logon iniciar.")

    SAPApp = SAPGui.GetScriptingEngine
    if SAPApp.Children.Count == 0:
        registrar_log("Abrindo conexão com a instância SAP (SAP_CONNECTION_NAME)...")
        SAPCon = SAPApp.OpenConnection("SAP_CONNECTION_NAME", True)
    else:
        SAPCon = SAPApp.Children(0)
    
    time.sleep(3)
    try:
        session = SAPCon.Children(0)
        session.findById("wnd[0]").maximize()
    except Exception:
        time.sleep(2)
        session = SAPCon.Children(0)
        session.findById("wnd[0]").maximize()

    return session

def encontrar_tabela_recursiva(elemento):
    if elemento.Type in ["GuiTableControl", "GuiGridView"]:
        return elemento
    try:
        for filho in elemento.Children:
            res = encontrar_tabela_recursiva(filho)
            if res: return res
    except:
        pass
    return None

def gerar_nome_arquivo_seguro(pasta, nome_base):
    base, ext = os.path.splitext(nome_base)
    caminho = os.path.join(pasta, nome_base)
    contador = 1
    
    while True:
        try:
            if os.path.exists(caminho):
                # Tenta abrir o arquivo para ver se não está trancado
                with open(caminho, 'a') as f:
                    pass
            return caminho
        except PermissionError:
            caminho = os.path.join(pasta, f"{base} ({contador}){ext}")
            contador += 1
        except Exception:
            # Qualquer outro erro, retorna com o nome seguro de contador pra tentar
            caminho = os.path.join(pasta, f"{base} ({contador}){ext}")
            contador += 1


_resultados_lock = threading.Lock()
_pesos_lock      = threading.Lock()

def _worker_wrapper(sess_idx, chunk_skus, centro_selecionado, resultados_shared, pesos_shared, barreira):
    import pythoncom as _pcom
    import win32com.client as win32
    _pcom.CoInitialize()
    try:
        SAPGui = win32.GetObject("SAPGUI")
        SAPApp = SAPGui.GetScriptingEngine
        SAPCon = SAPApp.Children(0)
        session = SAPCon.Children(sess_idx)
        session.findById("wnd[0]").maximize()

        # FASE 1: C203
        _worker_c203(session, sess_idx, chunk_skus, centro_selecionado, resultados_shared)

        # Sincroniza todas as threads antes de ir para a MM03
        try:
            barreira.wait()
        except:
            pass
            
        if _stop_flag.is_set():
            return
            
        # FASE 2: MM03
        chunk_mm03 = list(dict.fromkeys([str(s).strip() for s in chunk_skus if str(s).strip()]))
        _worker_mm03(session, sess_idx, chunk_mm03, pesos_shared)

    except Exception as e:
        registrar_log(f"[ERRO T-{sess_idx}] {e}", "erro")
        barreira.abort()
    finally:
        _pcom.CoUninitialize()

def _worker_c203(session, sess_idx, chunk_skus, centro_selecionado, resultados_shared):
    """Worker de C203 para uma fatia de SKUs em sua própria sessão SAP."""
    try:

        for sku in chunk_skus:
            if _stop_flag.is_set():
                break

            sku = str(sku).strip()
            if not sku:
                continue

            registrar_log(f"\n[Thread-{sess_idx}] SKU: {sku}", "destaque")
            session.findById("wnd[0]/tbar[0]/okcd").text = "/nC203"
            session.findById("wnd[0]").sendVKey(0)
            session.findById("wnd[0]/usr/ctxtRC271-PLNNR").text = ""
            session.findById("wnd[0]/usr/txtRC271-PLNAL").text = ""
            session.findById("wnd[0]/usr/ctxtRC27M-MATNR").text = sku
            session.findById("wnd[0]/usr/ctxtRC27M-WERKS").text = str(centro_selecionado)
            session.findById("wnd[0]/usr/ctxtRC27M-WERKS").setFocus()
            session.findById("wnd[0]").sendVKey(0)
            time.sleep(1)

            versoes_z = []
            try:
                if session.Children.Count > 1 and session.Children(1).Name == "wnd[1]":
                    popup = session.Children(1)
                    tabela_popup = encontrar_tabela_recursiva(popup)
                    if tabela_popup:
                        for i in range(tabela_popup.VisibleRowCount):
                            try:
                                caminho_celula = f"wnd[1]/usr/tblSAPLCPSLTCTRL_2110/txtPLKO-PLNAL[1,{i}]"
                                valor_z = str(session.findById(caminho_celula).text).strip()
                                if "Z" in valor_z.upper():
                                    versoes_z.append({"indice": i, "codigo": valor_z})
                            except:
                                pass
                    session.findById("wnd[1]").sendVKey(12)
            except:
                pass

            if not versoes_z:
                versoes_z.append({"indice": None, "codigo": "ÚNICA"})

            global _total_zs_encontrados
            with _resultados_lock:
                _total_zs_encontrados += len(versoes_z)
            registrar_log(f"[Thread-{sess_idx}] Versões Z: {len(versoes_z)}")

            for item_z in versoes_z:
                if _stop_flag.is_set():
                    break
                idx_linha = item_z["indice"]
                cod_z = item_z["codigo"]
                registrar_log(f"[Thread-{sess_idx}] Roteiro: [{cod_z}]")

                session.findById("wnd[0]/tbar[0]/okcd").text = "/nC203"
                session.findById("wnd[0]").sendVKey(0)
                session.findById("wnd[0]/usr/ctxtRC271-PLNNR").text = ""
                session.findById("wnd[0]/usr/txtRC271-PLNAL").text = ""
                session.findById("wnd[0]/usr/ctxtRC27M-MATNR").text = sku
                session.findById("wnd[0]/usr/ctxtRC27M-WERKS").text = str(centro_selecionado)
                session.findById("wnd[0]").sendVKey(0)
                time.sleep(1)

                if idx_linha is not None:
                    if _stop_flag.is_set():
                        break
                    caminho_celula = f"wnd[1]/usr/tblSAPLCPSLTCTRL_2110/txtPLKO-PLNAL[1,{idx_linha}]"
                    session.findById(caminho_celula).setFocus()
                    session.findById("wnd[1]").sendVKey(2)
                    time.sleep(1)

                desc_sku = "N/A"
                try:
                    desc_id = "wnd[0]/usr/subSUBSCREEN_RECIPE_GROUP:SAPLCPDI:4402/txtPLKOD-KTEXT"
                    session.findById(desc_id).setFocus()
                    desc_sku = str(session.findById(desc_id).text).strip()
                except:
                    pass

                recurso_original = "N/A"
                recurso_formatado = "N/A"
                try:
                    rec_id = "wnd[0]/usr/tabsTABSTRIP_RECIPE/tabpVOUE/ssubSUBSCREEN_RECIPE:SAPLCPDI:4401/tblSAPLCPDITCTRL_4401/ctxtPLPOD-ARBPL[4,0]"
                    session.findById(rec_id).setFocus()
                    recurso_original = str(session.findById(rec_id).text).strip()
                    recurso_formatado = DEPARA_LINHAS.get(recurso_original, recurso_original)
                except:
                    registrar_log(f"[Thread-{sess_idx}] Aviso: Linha de produção não identificada.", "aviso")

                unidade_base = calcular_unidade_base(recurso_formatado)

                try:
                    session.findById("wnd[0]/usr/tabsTABSTRIP_RECIPE/tabpOMBO").select()
                    time.sleep(0.5)
                    session.findById("wnd[0]/usr/tabsTABSTRIP_RECIPE/tabpOMBO/ssubSUBSCREEN_RECIPE:SAPLCMDI:4001/btnBUTTON_BOM").press()
                    time.sleep(1)
                except Exception as e:
                    registrar_log(f"[Thread-{sess_idx}] Erro ao navegar BOM [{cod_z}]: {e}", "erro")
                    continue

                versao_prod = "N/A"
                try:
                    versao_prod_id = "wnd[0]/usr/subSTLKOPF:SAPLCMDI:4500/ctxtMKAL-VERID"
                    versao_prod = str(session.findById(versao_prod_id).text).strip()
                except:
                    pass

                # BOM scrollbar loop
                tabela_id = "wnd[0]/usr/tabsTS_ITOV/tabpTCMA/ssubSUBPAGE:SAPLCSDI:0152/tblSAPLCSDITCMAT"
                try:
                    tabela = session.findById(tabela_id)
                    total_linhas      = tabela.rowCount
                    linhas_por_pagina = tabela.visibleRowCount
                except:
                    total_linhas      = 500
                    linhas_por_pagina = 20

                linhas_vazias   = 0
                max_linhas_vazias = 3
                parar_bom       = False
                ultima_abs_lida = -1
                ultima_pos_real = -1
                scroll_pos      = 0

                while not parar_bom:
                    if _stop_flag.is_set():
                        break
                    try:
                        session.findById(tabela_id).verticalScrollbar.position = scroll_pos
                        pos_real = session.findById(tabela_id).verticalScrollbar.position
                    except:
                        pos_real = scroll_pos

                    if pos_real == ultima_pos_real and ultima_pos_real >= 0:
                        break
                    ultima_pos_real = pos_real

                    for vis_idx in range(linhas_por_pagina):
                        linha_abs = pos_real + vis_idx
                        if linha_abs <= ultima_abs_lida:
                            continue
                        if linha_abs >= total_linhas:
                            parar_bom = True
                            break
                        if _stop_flag.is_set():
                            parar_bom = True
                            break
                        try:
                            insumo_id = f"{tabela_id}/ctxtRC29P-IDNRK[2,{vis_idx}]"
                            desc_id2  = f"{tabela_id}/txtRC29P-KTEXT[3,{vis_idx}]"
                            qtd_id    = f"{tabela_id}/txtRC29P-MENGE[4,{vis_idx}]"
                            um_id     = f"{tabela_id}/ctxtRC29P-MEINS[5,{vis_idx}]"
                            val_id    = f"{tabela_id}/ctxtRC29P-DATUV[9,{vis_idx}]"

                            insumo_val = str(session.findById(insumo_id).text).strip()
                            ultima_abs_lida = linha_abs

                            if not insumo_val:
                                linhas_vazias += 1
                                if linhas_vazias >= max_linhas_vazias:
                                    parar_bom = True
                                    break
                                continue

                            linhas_vazias = 0
                            insumo_limpo = insumo_val.replace("_", "").strip()
                            if not insumo_limpo:
                                continue

                            desc_val    = str(session.findById(desc_id2).text).strip()
                            qtd_val_raw = str(session.findById(qtd_id).text).strip()
                            um_val      = str(session.findById(um_id).text).strip()
                            qtd_val     = limpar_numero(qtd_val_raw)
                            qtd_p1      = qtd_val / 1000.0

                            if qtd_val == 0.0:
                                continue

                            registrar_log(f"  [T{sess_idx}] > {insumo_val} ({qtd_val} {um_val})")

                            validade_raw = ""
                            try:
                                validade_raw = str(session.findById(val_id).text).strip()
                            except:
                                pass
                            valido_desde = validade_raw
                            try:
                                if validade_raw and "." in validade_raw:
                                    valido_desde = datetime.strptime(validade_raw, "%d.%m.%Y")
                            except:
                                pass

                            tipo_insumo = classificar_insumo(insumo_val, desc_val)

                            with _resultados_lock:
                                resultados_shared.append({
                                    "Recurso": recurso_formatado,
                                    "Codigo Acabado": sku,
                                    "DescricaoAcabado": desc_sku,
                                    "Centro": int(centro_selecionado) if str(centro_selecionado).isdigit() else centro_selecionado,
                                    "Valido Desde": valido_desde,
                                    "Versão Produção": versao_prod,
                                    "Unidade Base": unidade_base,
                                    "Codigo Insumo": insumo_val,
                                    "Tipo Insumo": tipo_insumo,
                                    "Descricao Insumo": desc_val,
                                    "Unidade Medida": um_val,
                                    "Quantidade Componente": qtd_val,
                                    "Quantidade P/1": qtd_p1,
                                    "Peso Bruto": None,
                                    "Peso Liquido": None,
                                })
                        except:
                            parar_bom = True
                            break

                    scroll_pos = pos_real + linhas_por_pagina

            # SKU concluído neste worker
            global _fase_concluidos, _fase_inicio, _historico_tempos
            with _resultados_lock:
                _fase_concluidos += 1

    except Exception as e:
        registrar_log(f"[ERRO Thread-{sess_idx}] {e}", "erro")


def _worker_mm03(session, sess_idx, chunk_skus, pesos_shared):
    """Worker de MM03 para uma fatia de SKUs únicos."""
    try:
        for sku_mm03 in chunk_skus:
            if _stop_flag.is_set():
                break
            try:
                registrar_log(f"[MM03-T{sess_idx}] SKU: {sku_mm03}")
                session.findById("wnd[0]/tbar[0]/okcd").text = "/nMM03"
                session.findById("wnd[0]").sendVKey(0)
                time.sleep(0.5)
                session.findById("wnd[0]/usr/ctxtRMMG1-MATNR").text = str(sku_mm03)
                session.findById("wnd[0]/usr/ctxtRMMG1-MATNR").caretPosition = len(str(sku_mm03))
                session.findById("wnd[0]").sendVKey(0)
                time.sleep(0.8)

                try:
                    session.findById("wnd[1]/usr/tblSAPLMGMMTC_VIEW").getAbsoluteRow(0).selected = True
                    session.findById("wnd[1]/tbar[0]/btn[0]").press()
                    time.sleep(1)
                except:
                    pass

                peso_bruto = 0.0
                try:
                    pb_id = "wnd[0]/usr/tabsTABSPR1/tabpSP01/ssubTABFRA1:SAPLMGMM:2004/subSUB4:SAPLMGD1:2007/txtMARA-BRGEW"
                    session.findById(pb_id).setFocus()
                    peso_bruto = limpar_numero(str(session.findById(pb_id).text).strip())
                except:
                    pass

                peso_liquido = 0.0
                try:
                    pl_id = "wnd[0]/usr/tabsTABSPR1/tabpSP01/ssubTABFRA1:SAPLMGMM:2004/subSUB4:SAPLMGD1:2007/txtMARA-NTGEW"
                    session.findById(pl_id).setFocus()
                    peso_liquido = limpar_numero(str(session.findById(pl_id).text).strip())
                except:
                    pass

                with _pesos_lock:
                    pesos_shared[str(sku_mm03)] = {"Peso Bruto": peso_bruto, "Peso Liquido": peso_liquido}
                registrar_log(f"  [MM03-T{sess_idx}] PB: {peso_bruto} | PL: {peso_liquido}")

                global _fase_concluidos
                with _pesos_lock:
                    _fase_concluidos += 1

            except Exception as e_mm:
                registrar_log(f"  [MM03-T{sess_idx}] Aviso SKU {sku_mm03}: {e_mm}", "aviso")
                with _pesos_lock:
                    pesos_shared[str(sku_mm03)] = {"Peso Bruto": 0.0, "Peso Liquido": 0.0}
    except Exception as e:
        registrar_log(f"[ERRO MM03-T{sess_idx}] {e}", "erro")



def executar_extracao(lista_skus, pasta_destino, centro_selecionado, num_threads=1):
    global _fase_atual, _fase_inicio, _fase_concluidos, _fase_total, _historico_tempos, _total_zs_encontrados
    import pythoncom
    pythoncom.CoInitialize()
    session = None
    resultados = []
    _fase_concluidos = 0
    _historico_tempos = []
    _total_zs_encontrados = 0
    
    try:
        marcar_etapa("Conexão SAP", "ativo")
        set_status("Conectando ao SAP...")
        session = conectar_sap_ou_abrir()
        marcar_etapa("Conexão SAP", "concluido")

        # ── Configura fase C203 ───────────────────────────────────────
        _fase_atual      = "C203 — Receitas Mestres"
        _fase_inicio     = datetime.now()
        _fase_total      = len(lista_skus)
        _fase_concluidos = 0
        marcar_etapa("Extração SAP (C203)", "ativo")

        # Obtém a referência da conexão SAP para abrir sessões adicionais
        import win32com.client as win32
        SAPGui = win32.GetObject("SAPGUI")
        SAPApp = SAPGui.GetScriptingEngine
        SAPCon = SAPApp.Children(0)

        # Limita num_threads ao número de SKUs disponíveis
        num_threads = max(1, min(num_threads, len(lista_skus)))

        if num_threads > 1:
            registrar_log(f"⚡ Abrindo {num_threads - 1} sessão(ões) SAP extra(s)...")
            session.findById("wnd[0]").maximize()
            for i in range(1, num_threads):
                session.createSession()
                # Aguarda a nova sessão aparecer
                for _ in range(60):
                    if SAPCon.Children.Count > i:
                        break
                    time.sleep(0.5)
                time.sleep(2)
            registrar_log(f"✅ {SAPCon.Children.Count} sessão(ões) SAP abertas.")

        if num_threads == 1:
            # Lógica sequencial direta sem threads para poupar popups do SAP
            _worker_c203(session, 0, lista_skus, centro_selecionado, resultados)
            if _stop_flag.is_set():
                raise Exception("Execução interrompida pelo usuário.")
            marcar_etapa("Extração SAP (C203)", "concluido")

            marcar_etapa("Extração SAP (MM03)", "ativo")
            set_status("Extraindo pesos na MM03...")
            
            skus_unicos = list(dict.fromkeys([str(s).strip() for s in lista_skus if str(s).strip()]))
            _fase_atual      = "MM03 — Pesos de Materiais"
            _fase_inicio     = datetime.now()
            _fase_total      = len(skus_unicos)
            _fase_concluidos = 0
            _historico_tempos.clear()
            
            pesos_por_sku = {}
            _worker_mm03(session, 0, skus_unicos, pesos_por_sku)
            if _stop_flag.is_set():
                raise Exception("Execução interrompida pelo usuário.")
            marcar_etapa("Extração SAP (MM03)", "concluido")

        else:
            # Divide os SKUs em chunks para cada thread
            import math
            tamanho_chunk = math.ceil(len(lista_skus) / num_threads)
            chunks = [lista_skus[i:i + tamanho_chunk] for i in range(0, len(lista_skus), tamanho_chunk)]

            registrar_log(f"🔀 Dividindo {len(lista_skus)} SKUs em {len(chunks)} chunk(s) de ~{tamanho_chunk}")

            pesos_por_sku = {}
            
            def _prepara_mm03():
                global _fase_atual, _fase_inicio, _fase_total, _fase_concluidos, _historico_tempos
                _fase_atual      = "MM03 — Pesos de Materiais"
                _fase_inicio     = datetime.now()
                _fase_total      = len(set([str(s).strip() for s in lista_skus if str(s).strip()]))
                _fase_concluidos = 0
                _historico_tempos.clear()
                marcar_etapa("Extração SAP (MM03)", "ativo")
                set_status("Extraindo pesos na MM03...")

            barreira_threads = threading.Barrier(len(chunks), action=_prepara_mm03)

            # Dispara workers unificados em paralelo (cada um em sua própria sessão SAP)
            threads = []
            for i, chunk in enumerate(chunks):
                t = threading.Thread(
                    target=_worker_wrapper,
                    args=(i, chunk, centro_selecionado, resultados, pesos_por_sku, barreira_threads),
                    daemon=True
                )
                threads.append(t)
                t.start()
                time.sleep(0.5)  # escalonamento para evitar race condition na abertura

            # Aguarda TODOS os workers terminarem ambas as fases
            for t in threads:
                t.join()

            if _stop_flag.is_set():
                raise Exception("Execução interrompida pelo usuário.")

            marcar_etapa("Extração SAP (C203)", "concluido")
            marcar_etapa("Extração SAP (MM03)", "concluido")

        # Injeta pesos nas linhas dos resultados
        for r in resultados:
            sk = str(r["Codigo Acabado"])
            pesos = pesos_por_sku.get(sk, {"Peso Bruto": None, "Peso Liquido": None})
            r["Peso Bruto"]   = pesos["Peso Bruto"]
            r["Peso Liquido"] = pesos["Peso Liquido"]

        marcar_etapa("Extração SAP (MM03)", "concluido")

        # ---------------------------------------------------------
        # Exportação Excel Formatado
        marcar_etapa("Tratamento & Exportação", "ativo")
        set_status("Gerando planilha formatada...")
        
        if resultados:
            df = pd.DataFrame(resultados)

            # Garante que colunas B (Codigo Acabado) e H (Codigo Insumo) continuem como texto
            # mas sem os zeros à esquerda (simulando "Texto para Colunas")
            df["Codigo Acabado"] = df["Codigo Acabado"].astype("string").str.lstrip("0")
            df["Codigo Insumo"]  = df["Codigo Insumo"].astype("string").str.lstrip("0")

            # Ordena primeiro pela coluna B (Codigo Acabado) usando ordenação natural (de A a Z/0 a 9),
            # depois pela coluna A (Recurso) de A a Z
            def _ordenacao_natural_key(col):
                if col.name == "Codigo Acabado":
                    return col.fillna("").astype(str).str.replace(r'(\d+)', lambda m: m.group(0).zfill(20), regex=True)
                return col

            df = df.sort_values(by=["Codigo Acabado", "Recurso"], ascending=[True, True], key=_ordenacao_natural_key)

            data_atual = datetime.now().strftime("%d.%m")
            nome_arquivo_base = f"Lista Técnica - {data_atual}.xlsx"
            caminho_salvar = gerar_nome_arquivo_seguro(pasta_destino, nome_arquivo_base)

            # ── Cria workbook e escreve dados + formatação em passo único ──
            # Isso evita o ciclo save→load que faz openpyxl converter "21053840" → int
            from openpyxl import Workbook as _Workbook
            wb = _Workbook()
            ws = wb.active
            ws.title = "Lista Técnica"

            # Referências de coluna (0-based no df, 1-based no openpyxl)
            COL_NAMES  = list(df.columns)
            n_cols     = len(COL_NAMES)

            COR_HEADER  = "1E3A5F"
            COR_LINHA1  = "DBEAFE"
            COR_LINHA2  = "FFFFFF"
            FONTE_HDR   = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
            FONTE_BODY  = Font(name="Segoe UI", size=9)
            ALIGN_CTR   = Alignment(horizontal="center", vertical="center")
            ALIGN_LEFT  = Alignment(horizontal="left",   vertical="center")
            ALIGN_RIGHT = Alignment(horizontal="right",  vertical="center")
            borda_thin  = Border(
                left=Side(style="thin", color="BFD7ED"),
                right=Side(style="thin", color="BFD7ED"),
                top=Side(style="thin", color="BFD7ED"),
                bottom=Side(style="thin", color="BFD7ED"),
            )

            COLUNAS_TEXTO = {2, 8}
            COLUNAS_QTD   = {12, 13, 14, 15}
            COLUNA_DATA   = 5
            COLUNA_CENTRO = 4

            # ── Cabeçalho ────────────────────────────────────────────────
            for c_idx, col_name in enumerate(COL_NAMES, start=1):
                cell = ws.cell(row=1, column=c_idx, value=col_name)
                cell.font      = FONTE_HDR
                cell.fill      = PatternFill("solid", fgColor=COR_HEADER)
                cell.alignment = ALIGN_CTR
                cell.border    = borda_thin

            # ── Dados com formatação por coluna ──────────────────────────
            for r_idx, (_, row_data) in enumerate(df.iterrows(), start=2):
                fill_cor = COR_LINHA1 if r_idx % 2 == 0 else COR_LINHA2
                fill     = PatternFill("solid", fgColor=fill_cor)

                for c_idx, val in enumerate(row_data, start=1):
                    cell        = ws.cell(row=r_idx, column=c_idx)
                    cell.fill   = fill
                    cell.border = borda_thin
                    cell.font   = FONTE_BODY


                    if c_idx in COLUNAS_TEXTO:
                        # Converte para str ANTES de atribuir ao cell.value
                        # Isso garante que o openpyxl grava type="str" no XML
                        str_val = str(val).strip() if val is not None else ""
                        cell.value         = str_val
                        cell.data_type     = "s"    # Força tipo string válido ('s') no XML
                        cell.number_format = "@"
                        cell.quotePrefix   = True   # suprime aviso verde do Excel
                        cell.alignment     = ALIGN_LEFT

                    elif c_idx == COLUNA_DATA:
                        cell.value         = val
                        cell.number_format = "DD/MM/YYYY"
                        cell.alignment     = ALIGN_CTR

                    elif c_idx in COLUNAS_QTD:
                        cell.value         = val
                        cell.number_format = "#,##0.000"
                        cell.alignment     = ALIGN_RIGHT

                    elif c_idx == COLUNA_CENTRO:
                        cell.value         = val
                        cell.number_format = "0"
                        cell.alignment     = ALIGN_CTR

                    else:
                        cell.value     = val
                        cell.alignment = ALIGN_LEFT

            n_rows = ws.max_row

            # ── Tabela estilo Excel ──────────────────────────────────────
            ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
            tabela_excel = Table(displayName="ListaTecnica", ref=ref)
            tabela_excel.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(tabela_excel)

            # ── Largura automática ───────────────────────────────────────
            for col in ws.columns:
                max_len    = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        val_str = str(cell.value) if cell.value is not None else ""
                        max_len = max(max_len, len(val_str))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

            ws.freeze_panes = "A2"
            wb.save(caminho_salvar)
            registrar_log(f"\n✅ Planilha formatada salva em:\n{caminho_salvar}", "info")
            _arquivo_gerado = caminho_salvar

        else:
            registrar_log("\nNenhum dado encontrado para gerar planilha.", "aviso")
            _arquivo_gerado = ""
            
        marcar_etapa("Tratamento & Exportação", "concluido")
        set_status("Processo finalizado!")
        set_btn("done")
        notificar_windows(
            "TecList Forge — Concluído ✅",
            f"Extração finalizada com {len(resultados)} linha(s). Clique para abrir a pasta.",
            arquivo=_arquivo_gerado,
        )

    except Exception as ex:
        marcar_etapa("Conexão SAP", "erro")
        registrar_log(f"ERRO CRÍTICO:\n{traceback.format_exc()}", "erro")
        set_status("Falha na execução.")
        set_btn("erro")
        notificar_windows(
            "TecList Forge — Falha ❌",
            "Ocorreu um erro crítico durante a extração. Verifique o terminal.",
        )
    finally:
        try:
            session.findById("wnd[0]/tbar[0]/okcd").text = "/n"
            session.findById("wnd[0]").sendVKey(0)
        except:
            pass


def notificar_windows(titulo: str, mensagem: str, arquivo: str = ""):
    """Exibe uma notificação nativa do Windows via PowerShell (Toast ou Balloon Tip)."""
    try:
        t = str(titulo or "").replace("\r", " ").replace("\n", " ").strip()
        m = str(mensagem or "").replace("\r", " ").replace("\n", " ").strip()
        t_xml = html.escape(t, quote=False)
        m_xml = html.escape(m, quote=False)

        abrir = ""
        if arquivo:
            arquivo_ps = str(arquivo).replace("'", "''")
            abrir = f"""
try {{
    Start-Process explorer.exe -ArgumentList "/select,`"{arquivo_ps}`""
}} catch {{}}
"""

        ps = f"""$ErrorActionPreference = 'SilentlyContinue'
$t = @'
{t_xml}
'@
$m = @'
{m_xml}
'@
$ok = $false
try {{
    $null = [Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType=WindowsRuntime]
    $null = [Windows.Data.Xml.Dom.XmlDocument, Windows.Data.Xml.Dom, ContentType=WindowsRuntime]
    $AppId = '{{1AC14E77-02E7-4E5D-B744-2EB1AE5198B7}}\\WindowsPowerShell\\v1.0\\powershell.exe'
    $xml = @"
<toast>
    <visual>
        <binding template="ToastGeneric">
            <text>$t</text>
            <text>$m</text>
        </binding>
    </visual>
</toast>
"@
    $doc = New-Object Windows.Data.Xml.Dom.XmlDocument
    $doc.LoadXml($xml)
    $toast = New-Object Windows.UI.Notifications.ToastNotification $doc
    $notifier = [Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier($AppId)
    $notifier.Show($toast)
    $ok = $true
    Start-Sleep -Seconds 6
}} catch {{ $ok = $false }}
if (-not $ok) {{
    try {{
        Add-Type -AssemblyName System.Windows.Forms
        Add-Type -AssemblyName System.Drawing
        $n = New-Object System.Windows.Forms.NotifyIcon
        $n.Icon = [System.Drawing.SystemIcons]::Information
        $n.BalloonTipTitle = $t
        $n.BalloonTipText = $m
        $n.BalloonTipIcon = [System.Windows.Forms.ToolTipIcon]::Info
        $n.Visible = $true
        $n.ShowBalloonTip(9000)
        [System.Windows.Forms.Application]::DoEvents()
        Start-Sleep -Seconds 10
        $n.Visible = $false
        $n.Dispose()
    }} catch {{}}
}}
{abrir}
Remove-Item $MyInvocation.MyCommand.Path -Force -ErrorAction SilentlyContinue
"""
        fd, path = tempfile.mkstemp(suffix=".ps1", prefix="script_")
        os.close(fd)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            f.write(ps)

        flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        subprocess.Popen(
            ["powershell.exe", "-NoLogo", "-NoProfile", "-NonInteractive",
             "-WindowStyle", "Hidden", "-ExecutionPolicy", "Bypass", "-File", path],
            creationflags=flags,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except Exception:
        pass


root = tk.Tk()
root.title("TecList Forge — ACME Corp")
root.geometry("1150x800")
root.resizable(True, True)
root.configure(bg=BG_ROOT)

style = ttk.Style()
style.theme_use("clam")
style.configure("TCombobox",
    fieldbackground="#111827",
    background="#1f2937",
    foreground=TEXT_PRIMARY,
    selectbackground=ACCENT,
    selectforeground=TEXT_PRIMARY,
    bordercolor="#1f2937",
    arrowcolor=TEXT_PRIMARY,
)
style.map("TCombobox",
    fieldbackground=[("readonly", "#111827")],
    foreground=[("readonly", TEXT_PRIMARY)],
)

sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
root.geometry(f"1150x800+{(sw-1150)//2}+{(sh-800)//2}")

header = tk.Frame(root, bg=BG_HEADER, height=68)
header.pack(fill=tk.X)
header.pack_propagate(False)

tk.Frame(header, bg=ACCENT, width=4).pack(side=tk.LEFT, fill=tk.Y)
tk.Label(header, text="TECLIST FORGE", font=FNT_TITLE, fg=TEXT_PRIMARY, bg=BG_HEADER, padx=16).pack(side=tk.LEFT, pady=14)
tk.Label(header, text="Extração e Tratamento (BOM) — C203", font=FNT_SUB, fg=TEXT_MUTED, bg=BG_HEADER).pack(side=tk.LEFT, pady=14)

lbl_hora = tk.Label(header, text="", font=FNT_STATUS, fg=TEXT_MUTED, bg=BG_HEADER, padx=16)
lbl_hora.pack(side=tk.RIGHT, pady=14)

def _tick():
    lbl_hora.config(text=datetime.now().strftime("%d/%m/%Y  %H:%M:%S"))
    root.after(1000, _tick)
_tick()

cfg_inicial = carregar_config()

frame_superior = tk.Frame(root, bg=BG_ROOT)
frame_superior.pack(fill=tk.X, padx=20, pady=(15, 0))

frame_skus = tk.Frame(frame_superior, bg=BG_CARD, padx=20, pady=15)
frame_skus.pack(fill=tk.X)

tk.Label(frame_skus, text="COLE OS SKUS (Um por linha):", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD).pack(anchor="w")
txt_skus = tk.Text(frame_skus, height=4, font=("Segoe UI", 10), bg="#111827", fg=TEXT_PRIMARY, insertbackground=ACCENT, bd=0, padx=8, pady=8)
txt_skus.pack(fill=tk.X, pady=(5, 0))
txt_skus.insert(tk.END, cfg_inicial.get("ultimos_skus", ""))

frame_centro = tk.Frame(frame_skus, bg=BG_CARD)
frame_centro.pack(fill=tk.X, pady=(15, 0))

tk.Label(frame_centro, text="CENTRO / FILIAL:", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD).pack(side=tk.LEFT)

CENTROS = [
    "1000 - Plant-A",
    "2000 - Plant-B",
    "3000 - Plant-C",
]
var_centro = tk.StringVar()
cmb_centro = ttk.Combobox(frame_centro, textvariable=var_centro, values=CENTROS, state="readonly", width=25, font=("Segoe UI", 9))
cmb_centro.pack(side=tk.LEFT, padx=(10, 0))

# Seleciona o valor pré-salvo ou o default (1000 - Plant-A)
centro_salvo = cfg_inicial.get("centro", "1000 - Plant-A")
if centro_salvo in CENTROS:
    cmb_centro.set(centro_salvo)
else:
    cmb_centro.set("1000 - Plant-A")

tk.Label(frame_centro, text="THREADS:", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD).pack(side=tk.LEFT, padx=(25, 0))

var_threads = tk.StringVar(value=cfg_inicial.get("num_threads", "1"))
cmb_threads = ttk.Combobox(frame_centro, textvariable=var_threads, values=["1", "2", "3", "4"], state="readonly", width=4, font=("Segoe UI", 9))
cmb_threads.pack(side=tk.LEFT, padx=(5, 0))

threads_salvo = cfg_inicial.get("num_threads", "1")
if threads_salvo in ["1", "2", "3", "4"]:
    cmb_threads.set(threads_salvo)
else:
    cmb_threads.set("1")

frame_pasta = tk.Frame(frame_skus, bg=BG_CARD)
frame_pasta.pack(fill=tk.X, pady=(15, 0))

tk.Label(frame_pasta, text="PASTA DE SAÍDA:", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD).pack(side=tk.LEFT)

var_pasta = tk.StringVar(value=cfg_inicial.get("pasta_destino", os.path.dirname(os.path.abspath(__file__))))
if not var_pasta.get(): var_pasta.set(os.path.dirname(os.path.abspath(__file__)))

lbl_pasta = tk.Label(frame_pasta, textvariable=var_pasta, font= FNT_LABEL, fg=TEXT_MUTED, bg=BG_CARD)
lbl_pasta.pack(side=tk.LEFT, padx=(10, 10))

def selecionar_pasta():
    p = filedialog.askdirectory(initialdir=var_pasta.get(), title="Selecione onde salvar o Excel")
    if p:
        var_pasta.set(p)
        salvar_config({"pasta_destino": p})

btn_pasta = tk.Button(frame_pasta, text="MUDAR", font=("Segoe UI", 8, "bold"), fg=TEXT_PRIMARY, bg=ACCENT_DIM, relief=tk.FLAT, cursor="hand2", command=selecionar_pasta)
btn_pasta.pack(side=tk.RIGHT)


frame_db = tk.Frame(frame_skus, bg=BG_CARD)
frame_db.pack(fill=tk.X, pady=(15, 0))

tk.Label(frame_db, text="BASE DE INSUMOS:", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD).pack(side=tk.LEFT)

lbl_db_info = tk.Label(frame_db, text=f"v{DB_METADATA.get('versao', '1.0')} ({DB_METADATA.get('data_atualizacao', 'Fábrica')})", font=FNT_LABEL, fg=TEXT_MUTED, bg=BG_CARD)
lbl_db_info.pack(side=tk.LEFT, padx=(10, 10))

def exportar_base():
    try:
        caminho = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="Base_Insumos_Atualizada.xlsx", title="Exportar Base de Insumos")
        if not caminho: return
        
        # Para UX, tentamos parear Códigos e Descrições na mesma linha com base no Tipo
        agrupado_por_tipo = {}
        for chave, tipo in DB_INSUMOS.items():
            if tipo not in agrupado_por_tipo:
                agrupado_por_tipo[tipo] = {"codigos": [], "descricoes": []}
            
            k_str = str(chave).strip()
            is_codigo = k_str.isdigit() or (len(k_str) <= 12 and " " not in k_str and "," not in k_str)
            
            if is_codigo:
                agrupado_por_tipo[tipo]["codigos"].append(k_str)
            else:
                agrupado_por_tipo[tipo]["descricoes"].append(k_str)
                
        linhas = []
        for tipo, dados in agrupado_por_tipo.items():
            cods = sorted(dados["codigos"])
            descs = sorted(dados["descricoes"])
            
            max_len = max(len(cods), len(descs))
            for i in range(max_len):
                c = cods[i] if i < len(cods) else ""
                d = descs[i] if i < len(descs) else ""
                linhas.append({
                    "Codigo Insumo": c,
                    "Descricao Insumo": d,
                    "Tipo Insumo": tipo
                })
        
        linhas = sorted(linhas, key=lambda x: (x["Tipo Insumo"], x["Codigo Insumo"]))
        
        df = pd.DataFrame(linhas, columns=["Codigo Insumo", "Descricao Insumo", "Tipo Insumo"])
        df.to_excel(caminho, index=False)
        
        # ── Formatação com openpyxl ──
        try:
            wb = load_workbook(caminho)
            ws = wb.active
            
            max_row = ws.max_row
            max_col = ws.max_column
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            
            tabela = Table(displayName="BaseInsumos", ref=ref)
            estilo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            tabela.tableStyleInfo = estilo
            ws.add_table(tabela)
            
            # Centralizar texto nas colunas A e C
            align_center = Alignment(horizontal='center', vertical='center')
            align_left = Alignment(horizontal='left', vertical='center')
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=1).alignment = align_center
                ws.cell(row=row, column=2).alignment = align_left
                ws.cell(row=row, column=3).alignment = align_center
                
            # Ajustar larguras e fontes
            ws.column_dimensions['A'].width = 22
            ws.column_dimensions['B'].width = 65
            ws.column_dimensions['C'].width = 30
            
            ws.freeze_panes = "A2"
            wb.save(caminho)
        except Exception as ex:
            pass # Se openpyxl falhar, o raw pandas ja esta salvo
            
        messagebox.showinfo("Sucesso", f"Base exportada com formatação para:\n{caminho}")
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao exportar:\n{e}")

def sincronizar_base():
    try:
        caminho = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")], title="Selecione a Base Atualizada")
        if not caminho: return
        
        df = pd.read_excel(caminho)
        # Tenta achar colunas
        colunas = df.columns.tolist()
        if len(colunas) < 2:
            messagebox.showerror("Erro", "A planilha precisa ter pelo menos 2 colunas (Código/Descrição e Tipo).")
            return
            
        novos = 0
        atualizados = 0
        
        for index, row in df.iterrows():
            # A planilha pode ter 2 ou 3 colunas. Se tiver 3, a coluna 2 (index 2) é o Tipo.
            if len(colunas) >= 3:
                cod = str(row.iloc[0]).strip()
                desc = str(row.iloc[1]).strip()
                tipo = str(row.iloc[2]).strip()
                
                chaves = []
                if cod and str(cod).lower() != 'nan': chaves.append(cod)
                if desc and str(desc).lower() != 'nan': chaves.append(desc)
            else:
                chave = str(row.iloc[0]).strip()
                tipo = str(row.iloc[1]).strip()
                chaves = [chave] if chave and str(chave).lower() != 'nan' else []

            for c in chaves:
                if c in DB_INSUMOS:
                    if DB_INSUMOS[c] != tipo:
                        DB_INSUMOS[c] = tipo
                        atualizados += 1
                else:
                    DB_INSUMOS[c] = tipo
                    novos += 1
                
        if novos > 0 or atualizados > 0:
            # Atualiza metadata
            versao_atual = str(DB_METADATA.get("versao", "1.0"))
            try:
                nova_versao = str(float(versao_atual.replace("v", "")) + 0.1)[:3]
            except:
                nova_versao = "2.0"
                
            DB_METADATA["versao"] = nova_versao
            DB_METADATA["data_atualizacao"] = datetime.now().strftime("%d/%m/%Y %H:%M")
            
            DB_COMPLETO["data"] = DB_INSUMOS
            DB_COMPLETO["_metadata"] = DB_METADATA
            
            salvar_db_hibrido(DB_COMPLETO)
            
            lbl_db_info.config(text=f"v{DB_METADATA['versao']} ({DB_METADATA['data_atualizacao']})")
            messagebox.showinfo("Sucesso", f"Base sincronizada!\nNovos inseridos: {novos}\nAtualizados: {atualizados}")
        else:
            messagebox.showinfo("Aviso", "Nenhuma alteração encontrada em relação à base atual.")
            
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao sincronizar:\n{e}")

btn_sync = tk.Button(frame_db, text="SINCRONIZAR", font=("Segoe UI", 8, "bold"), fg=TEXT_PRIMARY, bg=ACCENT_DIM, relief=tk.FLAT, cursor="hand2", command=sincronizar_base)
btn_sync.pack(side=tk.RIGHT)

btn_export = tk.Button(frame_db, text="EXPORTAR", font=("Segoe UI", 8, "bold"), fg=TEXT_PRIMARY, bg=BG_HEADER, relief=tk.FLAT, cursor="hand2", command=exportar_base)
btn_export.pack(side=tk.RIGHT, padx=(0, 5))

ETAPAS = ["Conexão SAP", "Extração SAP (C203)", "Extração SAP (MM03)", "Tratamento & Exportação"]
frame_pipe = tk.Frame(root, bg=BG_ROOT, pady=10)
frame_pipe.pack(fill=tk.X, padx=20)

step_widgets = {}
for i, nome in enumerate(ETAPAS):
    col = tk.Frame(frame_pipe, bg=BG_ROOT)
    col.pack(side=tk.LEFT, expand=True)
    dot = tk.Label(col, text="●", font=("Segoe UI", 10), fg=TEXT_MUTED, bg=BG_ROOT)
    dot.pack()
    lbl = tk.Label(col, text=nome, font=FNT_STATUS, fg=TEXT_MUTED, bg=BG_ROOT)
    lbl.pack()
    step_widgets[nome] = (dot, lbl)
    if i < len(ETAPAS) - 1:
        tk.Label(frame_pipe, text="──", font=("Segoe UI", 9), fg=TEXT_MUTED, bg=BG_ROOT).pack(side=tk.LEFT, expand=True)

def marcar_etapa(nome, estado="ativo"):
    cores = {
        "ativo":     (ACCENT,     TEXT_PRIMARY),
        "concluido": (ACCENT_DIM, TEXT_MUTED),
        "erro":      (TEXT_ERROR, TEXT_ERROR),
    }
    cor_dot, cor_lbl = cores.get(estado, (TEXT_MUTED, TEXT_MUTED))
    dot, lbl = step_widgets[nome]
    root.after(0, lambda: (dot.config(fg=cor_dot), lbl.config(fg=cor_lbl)))



# Footer com barra de tempo + botões
footer = tk.Frame(root, bg=BG_FOOTER)
footer.pack(fill=tk.X, side=tk.BOTTOM)

# Linha de tempo (elapsed / ETA)
frame_timer = tk.Frame(footer, bg=BG_FOOTER)
frame_timer.pack(fill=tk.X, padx=12, pady=(8, 0))

var_timer = tk.StringVar(value="")
lbl_timer = tk.Label(frame_timer, textvariable=var_timer, font=("Cascadia Code", 8),
                     fg="#4b5563", bg=BG_FOOTER, anchor="w")
lbl_timer.pack(side=tk.LEFT)

# Linha de status + botões
frame_footer_bottom = tk.Frame(footer, bg=BG_FOOTER, height=48)
frame_footer_bottom.pack(fill=tk.X)
frame_footer_bottom.pack_propagate(False)

status_dot = tk.Label(frame_footer_bottom, text="●", font=("Segoe UI", 9), fg=ACCENT, bg=BG_FOOTER)
status_dot.pack(side=tk.LEFT, padx=(12, 4), pady=12)

status_var = tk.StringVar(value="Pronto para uso.")
tk.Label(frame_footer_bottom, textvariable=status_var, font=FNT_STATUS,
         fg=TEXT_MUTED, bg=BG_FOOTER).pack(side=tk.LEFT)

def _atualizar_timer():
    """Atualiza o label de elapsed/ETA a cada segundo enquanto há execução."""
    global _timer_ativo
    if not _timer_ativo:
        return
    if _tempo_inicio is None:
        root.after(1000, _atualizar_timer)
        return

    # ── Tempo total decorrido ────────────────────────────────────────
    elapsed = (datetime.now() - _tempo_inicio).total_seconds()
    def _fmt(s):
        h = int(s // 3600); m = int((s % 3600) // 60); ss = int(s % 60)
        return f"{h:02d}:{m:02d}:{ss:02d}"
    elapsed_str = _fmt(elapsed)

    # ── ETA e média determinística da fase atual ────────────────────
    fase       = _fase_atual or "Iniciando..."
    concluidos = _fase_concluidos
    total      = _fase_total

    if total > 0 and _fase_inicio:
        restantes = total - concluidos

        if "C203" in fase:
            # Regra: numero de Z x 4 segundos por SKU
            if concluidos > 0:
                media_zs = _total_zs_encontrados / concluidos
            else:
                media_zs = 1.0  # Estimativa inicial antes do 1º SKU concluir
                
            tempo_previsto_por_sku = media_zs * 4.0
            eta_seg = restantes * tempo_previsto_por_sku
            
        elif "MM03" in fase:
            # Regra: 3 segundos por SKU
            tempo_previsto_por_sku = 3.0
            eta_seg = restantes * tempo_previsto_por_sku
        else:
            tempo_previsto_por_sku = 0
            eta_seg = 0

        # Formata média
        if tempo_previsto_por_sku >= 60:
            media_str = f"{tempo_previsto_por_sku/60:.1f} min/SKU"
        else:
            media_str = f"{tempo_previsto_por_sku:.1f} s/SKU"

        eta_str = _fmt(eta_seg) if eta_seg > 0 else "concluído"

        var_timer.set(
            f"⏱ Decorrido: {elapsed_str}  │  🔄 {fase} [{concluidos}/{total}]  │  ⌀ {media_str}  │  ⏳ Término: {eta_str}"
        )
    else:
        msg_fase = f" │   {fase}" if fase else ""
        var_timer.set(f"⏱ Decorrido: {elapsed_str}{msg_fase}   │   Calculando previsão...")

    lbl_timer.config(fg="#6b7280")
    root.after(1000, _atualizar_timer)

def click_parar():
    """Sinaliza o cancelamento e ajusta a UI."""
    _stop_flag.set()
    btn_parar.config(state=tk.DISABLED, text="Parando...")
    set_status("⛔ Cancelamento solicitado...", TEXT_WARN)

def click_executar():
    global _tempo_inicio, _fase_atual, _fase_inicio, _fase_concluidos, _fase_total, _historico_tempos, _timer_ativo
    raw_text = txt_skus.get("1.0", tk.END)
    
    pasta_destino = var_pasta.get()

    centro_com_nome = cmb_centro.get().strip()
    centro_num = centro_com_nome.split(" ")[0] if " " in centro_com_nome else centro_com_nome

    num_threads = int(var_threads.get() or "1")

    salvar_config({"ultimos_skus": raw_text, "pasta_destino": pasta_destino, "centro": centro_com_nome, "num_threads": str(num_threads)})
    
    lista_skus = [s.strip() for s in raw_text.split("\n") if s.strip()]
    if not lista_skus:
        messagebox.showwarning("Aviso", "Por favor, insira ao menos um SKU para pesquisar.")
        return

    # Reseta estado
    _stop_flag.clear()
    _tempo_inicio    = datetime.now()
    _fase_atual      = "Conectando ao SAP..."
    _fase_inicio     = None
    _fase_concluidos = 0
    _fase_total      = 0
    _historico_tempos = []
    _total_zs_encontrados = 0
    _timer_ativo     = True

    set_btn("running")
    txt_log.config(state=tk.NORMAL)
    txt_log.delete("1.0", tk.END)
    txt_log.config(state=tk.DISABLED)
    var_timer.set("")

    for n in ETAPAS:
        marcar_etapa(n, "pendente")

    _atualizar_timer()
    t = threading.Thread(target=executar_extracao, args=(lista_skus, pasta_destino, centro_num, num_threads), daemon=True)
    t.start()

btn_parar = tk.Button(
    frame_footer_bottom, text="⏹  PARAR",
    font=FNT_BTN, fg=TEXT_PRIMARY, bg="#7f1d1d",
    activebackground=TEXT_ERROR, activeforeground=TEXT_PRIMARY,
    relief=tk.FLAT, padx=16, pady=6, cursor="hand2",
    command=click_parar, state=tk.DISABLED
)
btn_parar.pack(side=tk.RIGHT, padx=(4, 0), pady=6)

btn_executar = tk.Button(
    frame_footer_bottom, text="▶  EXECUTAR",
    font=FNT_BTN, fg=TEXT_PRIMARY, bg=ACCENT_DIM,
    activebackground=ACCENT, activeforeground=TEXT_PRIMARY,
    relief=tk.FLAT, padx=20, pady=6, cursor="hand2",
    command=click_executar
)
btn_executar.pack(side=tk.RIGHT, padx=(0, 16), pady=6)

def set_status(texto, cor=None):
    root.after(0, lambda: (
        status_var.set(texto),
        status_dot.config(fg=cor or ACCENT),
    ))

def set_btn(estado):
    global _timer_ativo
    configs_exec = {
        "normal":  dict(text="▶  EXECUTAR",          bg=ACCENT_DIM,  fg=TEXT_PRIMARY, state=tk.NORMAL),
        "running": dict(text="⏳  EXECUTANDO…",       bg=BG_CARD,     fg=TEXT_MUTED,   state=tk.DISABLED),
        "done":    dict(text="↺  EXECUTAR NOVAMENTE", bg=ACCENT,      fg=TEXT_PRIMARY, state=tk.NORMAL),
        "erro":    dict(text="↺  TENTAR NOVAMENTE",   bg=TEXT_ERROR,  fg=TEXT_PRIMARY, state=tk.NORMAL),
    }
    configs_parar = {
        "normal":  dict(state=tk.DISABLED),
        "running": dict(state=tk.NORMAL,   text="⏹  PARAR",   bg="#7f1d1d"),
        "done":    dict(state=tk.DISABLED, text="⏹  PARAR",   bg="#7f1d1d"),
        "erro":    dict(state=tk.DISABLED, text="⏹  PARAR",   bg="#7f1d1d"),
    }
    cfg_e = configs_exec.get(estado, configs_exec["normal"])
    cfg_p = configs_parar.get(estado, configs_parar["normal"])

    def _apply():
        btn_executar.config(**cfg_e)
        btn_parar.config(**cfg_p)
        if estado in ("done", "erro"):
            global _timer_ativo
            _timer_ativo = False
            if _tempo_inicio:
                elapsed = (datetime.now() - _tempo_inicio).total_seconds()
                h = int(elapsed // 3600); m = int((elapsed % 3600) // 60); ss = int(elapsed % 60)
                elapsed_str = f"{h:02d}:{m:02d}:{ss:02d}"
                if estado == "done":
                    var_timer.set(f"⏱ Tempo total: {elapsed_str}   │   ✓ Execução concluída com sucesso.")
                else:
                    var_timer.set(f"⏱ Tempo decorrido: {elapsed_str}   │   ✗ Falha na execução.")
    root.after(0, _apply)


frame_log_outer = tk.Frame(root, bg=BG_CARD, padx=2, pady=2)
frame_log_outer.pack(fill=tk.BOTH, expand=True, padx=20, pady=(10, 15))

frame_log_inner = tk.Frame(frame_log_outer, bg=BG_CARD)
frame_log_inner.pack(fill=tk.BOTH, expand=True)

tk.Label(frame_log_inner, text="  ● TERMINAL DE EVENTOS", font=FNT_LABEL, fg=ACCENT, bg=BG_CARD, anchor="w", pady=6).pack(fill=tk.X, padx=8)
tk.Frame(frame_log_inner, bg=ACCENT_DIM, height=1).pack(fill=tk.X, padx=8)

log_frame = tk.Frame(frame_log_inner, bg="#0b0f19")
log_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

txt_log = tk.Text(log_frame, wrap=tk.WORD, bg="#0b0f19", fg=TEXT_LOG, font=FNT_LOG, bd=0, highlightthickness=0, selectbackground=ACCENT_DIM, insertbackground=ACCENT, state=tk.DISABLED, padx=8, pady=6)
scrollbar = tk.Scrollbar(log_frame, command=txt_log.yview, bg="#1f2937", troughcolor="#0b0f19", width=8)
txt_log.configure(yscrollcommand=scrollbar.set)
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
txt_log.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

txt_log.tag_configure("info",     foreground=TEXT_LOG)
txt_log.tag_configure("aviso",    foreground=TEXT_WARN)
txt_log.tag_configure("erro",     foreground=TEXT_ERROR)
txt_log.tag_configure("ts",       foreground=TEXT_MUTED)
txt_log.tag_configure("destaque", foreground=TEXT_PRIMARY)

def registrar_log(msg, nivel="info"):
    def _inserir():
        ts = datetime.now().strftime("%H:%M:%S")
        txt_log.config(state=tk.NORMAL)
        txt_log.insert(tk.END, f"[{ts}]  ", "ts")
        txt_log.insert(tk.END, f"{msg}\n", nivel)
        txt_log.see(tk.END)
        txt_log.config(state=tk.DISABLED)
    root.after(0, _inserir)

if __name__ == "__main__":
    root.mainloop()